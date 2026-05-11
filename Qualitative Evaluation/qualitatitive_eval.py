import os
import random
import pandas as pd
import time
import subprocess
import shutil
import uuid
from types import SimpleNamespace
from docx import Document
from docx.shared import Inches
from llava.eval.run_llava import eval_model
from llava.mm_utils import get_model_name_from_path

prompt = "see the image and tell me what type of disease ,visible symptoms and possible solution?"
data_root = "/home/user/Downloads/agrogpt-20260218T141948Z-1-001/agrogpt/LLaVA/test/banana"
lora_checkpoint = "/home/user/Downloads/agrogpt-20260218T141948Z-1-001/agrogpt/LLaVA/checkpoints/llava-v1.5-7b-lora-banana-vlm_new_epoch_7"
base_model = "liuhaotian/llava-v1.5-7b"
OUTPUT_CSV = "banana_eval_results.csv"
OUTPUT_DOCX = "banana_report.docx"

TOTAL_SAMPLES=200
random.seed(42)


def get_gpu_stats():
    try:
        result=subprocess.check_output(
            [
                "nvidia-smi",
                "--query-gpu=memory.used,utilization.gpu",
                "--format=csv,noheader,nounits"
            ]
        ).decode("utf-8").strip()

        mem, util=result.split(", ")
        return int(mem),int(util)
    except:
        return -1,-1



def get_balanced_samples(root_dir,total_samples=100):
    class_folders=[
        f for f in os.listdir(root_dir)
        if os.path.isdir(os.path.join(root_dir,f))
    ]

    num_classes=len(class_folders)
    per_class=total_samples//num_classes

    all_samples=[]

    for cls in class_folders:
        cls_path=os.path.join(root_dir,cls)

        images=[
            img for img in os.listdir(cls_path)
            if img.lower().endswith((".jpg",".jpeg",".png"))
        ]

        selected=random.sample(images,min(per_class,len(images)))

        for img in selected:
            all_samples.append({
                "class":cls,
                "image_path":os.path.join(cls_path, img)
            })

    return all_samples


def run_model(model_path,model_base,image_path):
    temp_path=f"/tmp/{uuid.uuid4().hex}.jpg"
    shutil.copy(image_path,temp_path)

    args=SimpleNamespace(
        model_path=model_path,
        model_base=model_base,
        model_name=get_model_name_from_path(model_path),
        query=prompt,
        conv_mode=None,
        image_file=temp_path,
        sep=",",
        temperature=0,
        top_p=None,
        num_beams=1,
        max_new_tokens=512
    )

    import io
    import sys

    buffer=io.StringIO()
    sys_stdout=sys.stdout
    sys.stdout=buffer

    start_time = time.time()
    mem_before,util_before =get_gpu_stats()

    try:
        eval_model(args)
    finally:
        sys.stdout=sys_stdout

    mem_after,util_after=get_gpu_stats()
    end_time=time.time()

    try:
        os.remove(temp_path)
    except:
        pass

    return {
        "output":buffer.getvalue().strip(),
        "latency_ms":(end_time-start_time)*1000,
        "gpu_mem_mb":mem_after,
        "gpu_util":util_after
    }



def main():
    samples=get_balanced_samples(data_root,TOTAL_SAMPLES)
    results=[]
    doc=Document()
    doc.add_heading("Banana Disease Model Comparison Report",0)
    for i,sample in enumerate(samples):
        image_path=sample["image_path"]
        gt_class=sample["class"]

        print(f"[{i+1}/{len(samples)}] Processing image")
        finetuned_res=run_model(lora_checkpoint, base_model,image_path)
        base_res=run_model(base_model, None, image_path)
        results.append({
            "image_path":image_path,
            "ground_truth":gt_class,
            "finetuned_output":finetuned_res["output"],
            "base_output":base_res["output"]
        })
        doc.add_heading(f"Sample {i+1}",level=1)
        try:
            doc.add_picture(image_path,width=Inches(4))
        except:
            doc.add_paragraph("Image could not be loaded")

        doc.add_paragraph(f"Ground Truth: {gt_class}")
        doc.add_paragraph(f"Base Model Output:\n{base_res['output']}")
        doc.add_paragraph(f"Finetuned Model Output:\n{finetuned_res['output']}")
        doc.add_page_break()

    df=pd.DataFrame(results)
    df.to_csv(OUTPUT_CSV,index=False)
    doc.save(OUTPUT_DOCX)
    print(f"\nCSV saved → {OUTPUT_CSV}")
    print(f"Word report saved → {OUTPUT_DOCX}")

if __name__ == "__main__":
    main()










































































































































"""from llava.model.builder import load_pretrained_model
from llava.mm_utils import get_model_name_from_path
from llava.eval.run_llava import eval_model
from types import SimpleNamespace
from PIL import Image
import matplotlib.pyplot as plt
import torch

# -------------------------------
# Patch cache_position issue
# -------------------------------
from transformers.models.llama.modeling_llama import LlamaForCausalLM

if not getattr(LlamaForCausalLM, "_cachepos_patched", False):
    _orig_forward_llama = LlamaForCausalLM.forward
    def _patched_forward_llama(self, *args, **kwargs):
        kwargs.pop("cache_position", None)
        return _orig_forward_llama(self, *args, **kwargs)
    LlamaForCausalLM.forward = _patched_forward_llama
    LlamaForCausalLM._cachepos_patched = True


# -------------------------------
# SETTINGS
# -------------------------------
prompt = "What type of disease is shown in this image? Give the answer briefly."

image_file = "/home/user/Downloads/agrogpt-20260218T141948Z-1-001/agrogpt/LLaVA/test/banana/cordana/cordana-1-_jpeg.rf.ccd82d9ddd7a793dc89441410f34f4a7.jpg"

# IMPORTANT: use the correct checkpoint directory
lora_checkpoint = "/home/user/Downloads/agrogpt-20260218T141948Z-1-001/agrogpt/LLaVA/checkpoints/llava-v1.5-7b-lora-banana-vlm_new"

base_model = "liuhaotian/llava-v1.5-7b"


# -------------------------------
# Display image
# -------------------------------
img = Image.open(image_file)
plt.imshow(img)
plt.axis("off")
plt.title("Input Image")
plt.show()


# -------------------------------
# Evaluation Function
# -------------------------------
def run_eval(model_path, model_base=None, label="Model"):
    args = SimpleNamespace(
        model_path=model_path,
        model_base=model_base,
        model_name=get_model_name_from_path(model_path),
        query=prompt,
        conv_mode=None,
        image_file=image_file,
        sep=",",
        temperature=0,
        top_p=None,
        num_beams=1,
        max_new_tokens=128
    )

    print(f"\n========== {label} ==========")
    eval_model(args)


# -------------------------------
# 1️⃣ Finetuned LoRA Model
# -------------------------------
run_eval(
    model_path=lora_checkpoint,
    model_base=base_model,
    label="Finetuned LoRA Model"
)

# -------------------------------
# 2️⃣ Base Model
# -------------------------------
run_eval(
    model_path=base_model,
    model_base=None,
    label="Base Model"
)
"""
"""========== Finetuned LoRA Model ==========
/home/user/miniconda3/envs/llava/lib/python3.10/site-packages/huggingface_hub/file_download.py:942: FutureWarning: `resume_download` is deprecated and will be removed in version 1.0.0. Downloads always resume when possible. If you want to force a new download, use `force_download=True`.
  warnings.warn(
Loading LLaVA from base model...
Loading checkpoint shards:   0%|                                                                                                             | 0/2 [00:00<?, ?it/s]/home/user/miniconda3/envs/llava/lib/python3.10/site-packages/torch/_utils.py:831: UserWarning: TypedStorage is deprecated. It will be removed in the future and UntypedStorage will be the only storage class. This should only matter to you if you are using storages directly.  To access UntypedStorage directly, use tensor.untyped_storage() instead of tensor.storage()
  return self.fget.__get__(instance, owner)()
Loading checkpoint shards: 100%|█████████████████████████████████████████████████████████████████████████████████████████████████████| 2/2 [00:02<00:00,  1.32s/it]
Loading additional LLaVA weights...
Loading LoRA weights...
Merging LoRA weights...
Model is loaded...
/home/user/miniconda3/envs/llava/lib/python3.10/site-packages/transformers/generation/configuration_utils.py:392: UserWarning: `do_sample` is set to `False`. However, `temperature` is set to `0` -- this flag is only used in sample-based generation modes. You should set `do_sample=True` or unset `temperature`.
  warnings.warn(
/home/user/miniconda3/envs/llava/lib/python3.10/site-packages/transformers/generation/configuration_utils.py:397: UserWarning: `do_sample` is set to `False`. However, `top_p` is set to `None` -- this flag is only used in sample-based generation modes. You should set `do_sample=True` or unset `top_p`.
  warnings.warn(
The image shows a diseased cordana plant from the banana dataset, characterized by its discolored and withered leaves.

========== Base Model ==========
You are using a model of type llava to instantiate a model of type llava_llama. This is not supported for all configurations of models and can yield errors.
Loading checkpoint shards: 100%|█████████████████████████████████████████████████████████████████████████████████████████████████████| 2/2 [00:02<00:00,  1.49s/it]
The image shows a disease called bacterial leaf spot, which is characterized by brown spots on the leaves of a plant."""

"""
er, `top_p` is set to `None` -- this flag is only used in sample-based generation modes. You should set `do_sample=True` or unset `top_p`.
  warnings.warn(
The image shows a diseased cordana plant from the banana dataset.

========== Base Model ==========
You are using a model of type llava to instantiate a model of type llava_llama. This is not supported for all configurations of models and can yield errors.
Loading checkpoint shards: 100%|█████████████████████████████████████████████████████████████████████████████████████████████████████| 2/2 [00:03<00:00,  1.52s/it]
The image shows a disease called bacterial leaf spot, which is characterized by brown spots on the leaves of a plant."""